# -*- coding: utf-8 -*-
"""
Conversor de extrato digitalizado (OCR) para Excel + OFX - Jean Vieira
Ferramenta SEPARADA do Conversor_Extratos principal. So usa essa aqui
quando o PDF for uma imagem escaneada/foto/print sem texto (o conversor
principal ja avisa quando e esse o caso). Ve o aviso de responsabilidade
no console e na planilha antes de confiar em qualquer resultado - OCR
erra, mesmo quando o numero parece certo. Ver "Extrato digitalizado
(OCR)" no CONTEXTO_PROJETO.md pra entender as limitacoes conhecidas.

Uso: python conversor_ocr.py caminho_do_extrato.pdf [outro.pdf ...]
"""
import sys
import os
import glob
import io

import fitz
import pytesseract
from PIL import Image
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from bancos_ocr import linhas_por_posicao, DETECTORES_OCR, PARSERS_OCR
from converter import gerar_excel, NOMES_BANCO
from ofx_export import gerar_ofx

_PREENCHIMENTO_INCERTA = PatternFill(start_color='F4CCCC', end_color='F4CCCC', fill_type='solid')


def _destacar_linhas_incertas(caminho_xlsx):
    """Colore de vermelho claro as linhas marcadas [A VERIFICAR] - o
    `gerar_excel` do converter.py e generico (reusado tambem pelo
    conversor principal) e nao faz esse tipo de destaque condicional,
    entao aplica por cima aqui, so nesse arquivo, sem tocar em
    converter.py."""
    wb = load_workbook(caminho_xlsx)
    for nome_aba in wb.sheetnames:
        if nome_aba == 'Pendencias':
            continue
        ws = wb[nome_aba]
        for row in ws.iter_rows(min_row=2):
            obs = row[4].value if len(row) > 4 else None
            if obs and 'A VERIFICAR' in str(obs):
                for cel in row:
                    cel.fill = _PREENCHIMENTO_INCERTA
    wb.save(caminho_xlsx)

ESCALA_UPSCALE = 2  # ampliar a imagem antes do OCR melhora a leitura de C/D colado no valor
DPI_RENDER = 300


def _pasta_do_programa():
    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


def _configurar_tesseract():
    """Quando empacotado (.exe), o Tesseract vai embutido dentro do
    bundle do PyInstaller - aponta pytesseract pra ele em vez de
    depender de uma instalacao separada no PC de quem usa."""
    if getattr(sys, 'frozen', False):
        base = sys._MEIPASS
        tesseract_exe = os.path.join(base, 'tesseract', 'tesseract.exe')
        if os.path.exists(tesseract_exe):
            pytesseract.pytesseract.tesseract_cmd = tesseract_exe
            os.environ['TESSDATA_PREFIX'] = os.path.join(base, 'tesseract', 'tessdata')


def _extrair_linhas_pdf(caminho_pdf):
    """Rasteriza cada pagina do PDF e reconstroi as linhas de texto via
    OCR + posicao (ver bancos_ocr.linhas_por_posicao)."""
    doc = fitz.open(caminho_pdf)
    todas_linhas = []
    for page in doc:
        pix = page.get_pixmap(dpi=DPI_RENDER)
        img = Image.open(io.BytesIO(pix.tobytes('png')))
        if ESCALA_UPSCALE != 1:
            img = img.resize((img.width * ESCALA_UPSCALE, img.height * ESCALA_UPSCALE), Image.LANCZOS)
        todas_linhas.extend(linhas_por_posicao(img, escala=ESCALA_UPSCALE))
    doc.close()
    return todas_linhas


def identificar_banco_ocr(linhas):
    for nome_banco, detector in DETECTORES_OCR:
        if detector(linhas):
            return nome_banco
    return None


def processar_pdf_ocr(caminho_pdf, pasta_saida):
    nome_base = os.path.splitext(os.path.basename(caminho_pdf))[0]
    linhas = _extrair_linhas_pdf(caminho_pdf)
    banco = identificar_banco_ocr(linhas)

    if banco is None:
        avisos = [('(banco nao identificado - OCR)',
                    'O layout desse extrato digitalizado nao e reconhecido ainda. '
                    'Mande o PDF pra incluir suporte.')]
        caminho_saida = os.path.join(pasta_saida, f'{nome_base}.xlsx')
        gerar_excel({}, avisos, caminho_saida)
        return caminho_saida, 0, avisos, []

    parser = PARSERS_OCR[banco]
    transacoes, aviso = parser(linhas)
    avisos = [(banco, aviso)] if aviso else []

    caminho_saida = os.path.join(pasta_saida, f'{nome_base}.xlsx')
    gerar_excel({banco: transacoes}, avisos, caminho_saida)
    _destacar_linhas_incertas(caminho_saida)

    # transacoes com sinal incerto NUNCA entram no ofx - nao da pra inventar
    # credito ou debito num arquivo que alimenta o sistema contabil direto.
    transacoes_confiaveis = [t for t in transacoes if not t.get('incerta')]
    arquivos_ofx = []
    if transacoes_confiaveis:
        titulo_banco = NOMES_BANCO.get(banco, banco)
        caminho_ofx = os.path.join(pasta_saida, f'{nome_base}_{titulo_banco}.ofx')
        if gerar_ofx(transacoes_confiaveis, banco, caminho_ofx):
            arquivos_ofx.append(caminho_ofx)

    return caminho_saida, len(transacoes), avisos, arquivos_ofx


if __name__ == '__main__':
    _configurar_tesseract()
    pasta_programa = _pasta_do_programa()

    print('============================================')
    print('  Conversor OCR (extrato digitalizado) - Jean Vieira')
    print('============================================')
    print()
    print('ATENCAO: essa ferramenta le extrato por IMAGEM (OCR), nao por')
    print('texto nativo do PDF. Erros de leitura SAO ESPERADOS, mesmo em')
    print('numeros que parecem corretos. Toda linha gerada vem marcada')
    print('"Gerado por OCR" e precisa ser conferida contra o extrato')
    print('original antes de usar em qualquer lancamento contabil.')
    print('============================================')
    print()

    if len(sys.argv) > 1:
        arquivos = sys.argv[1:]
    else:
        arquivos = sorted(glob.glob(os.path.join(pasta_programa, '*.pdf')))

    if not arquivos:
        print('Nenhum arquivo PDF encontrado.')
        print(f'Coloque os extratos digitalizados nesta pasta ({pasta_programa})')
        print('e execute o programa novamente, ou arraste os PDFs em cima do .exe.')
        input('\nPressione Enter para sair...')
        sys.exit(0)

    total_arquivos = 0
    total_erros = 0
    for caminho in arquivos:
        pasta_saida = os.path.dirname(os.path.abspath(caminho))
        try:
            caminho_saida, total, avisos, ofx = processar_pdf_ocr(caminho, pasta_saida)
            print(f'{caminho} -> {caminho_saida} ({total} lancamento(s), '
                  f'{len(avisos)} aviso(s), {len(ofx)} ofx gerado(s))')
        except Exception as e:
            total_erros += 1
            print(f'ERRO ao processar {caminho}: {e}')
            continue
        total_arquivos += 1

    print()
    print('============================================')
    print(f'Concluido. {total_arquivos} arquivo(s) processado(s), {total_erros} erro(s).')
    print('Os arquivos .xlsx e .ofx foram gerados nesta mesma pasta.')
    print('LEMBRETE: confira os valores contra o extrato original antes de usar.')
    print('============================================')
    input('\nPressione Enter para sair...')
